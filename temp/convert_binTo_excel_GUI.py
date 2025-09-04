import shutil
import tkinter as tk
from tkinter import filedialog, messagebox
from datetime import datetime
import numpy as np
import pandas as pd
from openpyxl.reader.excel import load_workbook
import os

def bin_to_xlsx(bin_data, reference_xlsx, output_folder):
    """Process binary data and generate an Excel file"""
    pd.options.mode.chained_assignment = None
    date_time = datetime.now().strftime("%d_%m_%Y %H.%M.%S")
    ps_xlsx = os.path.join(output_folder, f"step_3_GLD_settings_{date_time}.xlsx")
    shutil.copy(reference_xlsx, ps_xlsx)

    wb = load_workbook(ps_xlsx, data_only=True)
    ws = wb['PATCH_SETTINGS']
    fields = []

    for rowNum in range(2, ws.max_row + 1):
        state = str(ws.cell(row=rowNum, column=4).value)
        if state == 'N/A':
            field_name = ws.cell(row=rowNum, column=5).value
        else:
            field_name = f"{ws.cell(row=rowNum, column=5).value} (State {str(round(ws.cell(row=rowNum, column=4).value))})"
        
        offset = int(ws.cell(row=rowNum, column=1).value)
        size = int(ws.cell(row=rowNum, column=2).value)
        end = offset + size
        int_data = np.frombuffer(bin_data[offset:end], dtype=f"<u{size}")
        ws.cell(row=rowNum, column=6).value = int_data[0]
        fields.append(field_name)

    wb.save(ps_xlsx)
    wb.close()

    messagebox.showinfo("Success", f"File saved:\n{ps_xlsx}")

def select_reference_file():
    """Open file dialog to select the GLD settings reference (.xlsx)"""
    file_path = filedialog.askopenfilename(
        title="Select GLD Reference File",
        filetypes=[("Excel Files", "*.xlsx")],
    )
    if file_path:
        reference_file_entry.delete(0, tk.END)
        reference_file_entry.insert(0, file_path)

def select_bin_file():
    """Open file dialog to select the GLD settings binary file (.bin)"""
    file_path = filedialog.askopenfilename(
        title="Select GLD Settings Binary File",
        filetypes=[("Binary Files", "*.bin")],
    )
    if file_path:
        bin_file_entry.delete(0, tk.END)
        bin_file_entry.insert(0, file_path)

def select_output_folder():
    """Open folder dialog to choose where to save the processed file"""
    folder_path = filedialog.askdirectory(title="Select Output Folder")
    if folder_path:
        output_folder_entry.delete(0, tk.END)
        output_folder_entry.insert(0, folder_path)

def process_files():
    """Validate inputs and process the selected files"""
    reference_xlsx = reference_file_entry.get()
    bin_file_path = bin_file_entry.get()
    output_folder = output_folder_entry.get()

    if not reference_xlsx or not bin_file_path or not output_folder:
        messagebox.showerror("Error", "Please select all files and output location.")
        return

    try:
        with open(bin_file_path, "rb") as f:
            bin_data = f.read()
        bin_to_xlsx(bin_data, reference_xlsx, output_folder)
    except Exception as e:
        messagebox.showerror("Error", f"Failed to process files:\n{e}")

# Create main GUI window
root = tk.Tk()
root.title("GLD Settings Processor")
root.geometry("500x350")

# Labels and Entry Fields
tk.Label(root, text="Select GLD Reference File (.xlsx):").pack(anchor="w", padx=10, pady=5)
reference_file_entry = tk.Entry(root, width=50)
reference_file_entry.pack(padx=10, pady=2)
tk.Button(root, text="Browse", command=select_reference_file).pack(pady=2)

tk.Label(root, text="Select GLD Settings Binary File (.bin):").pack(anchor="w", padx=10, pady=5)
bin_file_entry = tk.Entry(root, width=50)
bin_file_entry.pack(padx=10, pady=2)
tk.Button(root, text="Browse", command=select_bin_file).pack(pady=2)

tk.Label(root, text="Select Output Folder:").pack(anchor="w", padx=10, pady=5)
output_folder_entry = tk.Entry(root, width=50)
output_folder_entry.pack(padx=10, pady=2)
tk.Button(root, text="Browse", command=select_output_folder).pack(pady=2)

# Process Button
tk.Button(root, text="Process Files", command=process_files, bg="green", fg="white").pack(pady=10)

# Run the GUI loop
root.mainloop()
